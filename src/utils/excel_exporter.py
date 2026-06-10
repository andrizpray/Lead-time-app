from __future__ import annotations

from datetime import date
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    pass

_HEADER_FILL = PatternFill("solid", fgColor="3B82F6")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_BODY_FONT   = Font(name="Calibri", size=10)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT   = Alignment(horizontal="left",   vertical="center")


def _style_sheet(ws, n_cols: int) -> None:
    """Apply header style + auto-fit + borders to a worksheet."""
    # Header row (row 1)
    for cell in ws[1]:
        cell.fill   = _HEADER_FILL
        cell.font   = _HEADER_FONT
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER

    # Data rows
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font   = _BODY_FONT
            cell.border = _THIN_BORDER
            cell.alignment = _LEFT

    # Auto-fit column widths
    for col_idx in range(1, n_cols + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[letter].width = max(max_len + 4, 10)


def _fmt_lead_time(minutes: float) -> str:
    m = int(minutes)
    h, rem = divmod(m, 60)
    return f"{h}j {rem}m" if h else f"{rem} mnt"


# ─────────────────────────────────────────────────────────────────────────────
# Report dict → Excel (multi-sheet aware)
# ─────────────────────────────────────────────────────────────────────────────

_REPORT_COLUMNS: dict[str, list[tuple[str, str]]] = {
    "Harian": [
        ("Shift",               "shift"),
        ("Jumlah DO",           "count_do"),
        ("Total Tonase (kg)",  "total_tonase"),
        ("Avg Lead Time",       "avg_lead_time"),
    ],
    "Mingguan": [
        ("Minggu Ke",           "minggu_ke"),
        ("Tgl Mulai",           "tgl_mulai"),
        ("Tgl Akhir",           "tgl_akhir"),
        ("Total DO",            "total_do"),
        ("Total Tonase (kg)",  "total_tonase"),
        ("Avg Lead Time",       "avg_lead_time"),
        ("Shift 1",             "shift_1_count"),
        ("Shift 2",             "shift_2_count"),
        ("Shift 3",             "shift_3_count"),
    ],
    "Bulanan": [
        ("Customer",            "customer"),
        ("Jumlah DO",           "count"),
        ("Total Tonase (kg)",  "tonase"),
        ("Avg Lead Time",       "avg_lead_time"),
    ],
    "Per Customer": [
        ("Customer",            "customer"),
        ("Jumlah DO",           "count_do"),
        ("Total Tonase (kg)",  "total_tonase"),
        ("Avg Lead Time",       "avg_lead_time"),
    ],
    "Per Ekspedisi": [
        ("Ekspedisi",           "ekspedisi"),
        ("Jumlah DO",           "count_do"),
        ("Total Tonase (kg)",  "total_tonase"),
        ("Avg Lead Time",       "avg_lead_time"),
    ],
}


def _rows_from_report(data: dict, report_type: str) -> list[dict]:
    if report_type == "Harian":
        total = data.get("total", {})
        return [
            {"shift": "Shift 1", **data.get("shift_1", {})},
            {"shift": "Shift 2", **data.get("shift_2", {})},
            {"shift": "Shift 3", **data.get("shift_3", {})},
            {"shift": "TOTAL",   **total},
        ]
    if report_type == "Mingguan":
        return data if isinstance(data, list) else []
    if report_type == "Bulanan":
        return data.get("per_customer", [])
    if report_type in ("Per Customer", "Per Ekspedisi"):
        return data if isinstance(data, list) else []
    return []


def export_report_to_excel(data: dict | list, report_type: str, filepath: str) -> str:
    """Export report data dict/list to Excel (.xlsx)."""
    cols   = _REPORT_COLUMNS.get(report_type, [])
    rows   = _rows_from_report(data, report_type)
    headers = [c[0] for c in cols]
    keys    = [c[1] for c in cols]

    records = []
    for row in rows:
        rec = {}
        for h, k in zip(headers, keys):
            val = row.get(k, "")
            if isinstance(val, float):
                if "lead_time" in k or "avg_lt" in k:
                    val = _fmt_lead_time(val)
                else:
                    val = round(val, 2)
            elif isinstance(val, date):
                val = val.strftime("%d/%m/%Y")
            rec[h] = val
        records.append(rec)

    df = pd.DataFrame(records, columns=headers) if records else pd.DataFrame(columns=headers)

    sheet_name = report_type[:31]  # Excel sheet name max 31 chars
    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        _style_sheet(ws, len(headers))

    return filepath


# ─────────────────────────────────────────────────────────────────────────────
# DORecord list → Excel
# ─────────────────────────────────────────────────────────────────────────────

_DO_COLUMNS: list[tuple[str, str]] = [
    ("No",              None),
    ("Tanggal",         "tgl"),
    ("No DO",           "no_do"),
    ("No Shipment",     "no_shipment"),
    ("Customer",        "customer"),
    ("Kota/Kab",        "kota_kab"),
    ("Jenis",           "jenis"),
    ("Ekspedisi",       "ekspedisi"),
    ("Jenis Truk",      "jenis_truk"),
    ("Nomor FK",        "nomor_fk"),
    ("Loading Mulai",   "loading_mulai"),
    ("Loading Selesai", "loading_selesai"),
    ("Lead Time",       "lead_time_menit"),
    ("Tonase Roll",     "tonase_roll"),
    ("Tonase Sheet",    "tonase_sheet"),
    ("Tonase Total",    "tonase_total"),
    ("Shift",           "shift"),
]


def export_table_to_excel(records: list, filepath: str) -> str:
    """Export list of DORecord to Excel (.xlsx)."""
    headers = [c[0] for c in _DO_COLUMNS]
    rows = []

    for i, rec in enumerate(records, 1):
        row = {}
        for label, field in _DO_COLUMNS:
            if field is None:
                row[label] = i
                continue
            val = getattr(rec, field, None)
            if val is None:
                row[label] = ""
            elif field == "tgl":
                row[label] = val.strftime("%d/%m/%Y") if isinstance(val, date) else str(val)
            elif field in ("loading_mulai", "loading_selesai"):
                s = str(val)
                row[label] = s[:5] if len(s) >= 5 else s
            elif field == "lead_time_menit":
                row[label] = _fmt_lead_time(float(val))
            elif field in ("tonase_roll", "tonase_sheet", "tonase_total"):
                row[label] = round(float(val), 2)
            else:
                row[label] = str(val)
        rows.append(row)

    df = pd.DataFrame(rows, columns=headers) if rows else pd.DataFrame(columns=headers)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data DO")
        ws = writer.sheets["Data DO"]
        _style_sheet(ws, len(headers))

    return filepath
